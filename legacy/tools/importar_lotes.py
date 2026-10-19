#!/usr/bin/env python3
"""Genera inc/lots.php a partir de la planilla del cliente (lotes.xlsx, no versionada).

Uso:  python3 tools/importar_lotes.py [ruta/lotes.xlsx]

Solo exporta datos publicables: manzana, lote, superficie, tipo y estado.
Precios, cuentas corrientes y fideicomisos quedan fuera.
"""
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'lotes.xlsx'
OUT = ROOT / 'inc' / 'lots.php'


def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else None


def cumbres(ws):
    lots = []
    for r in range(7, ws.max_row + 1):
        mz, lote = ws.cell(r, 4).value, ws.cell(r, 5).value
        if mz is None or lote is None:
            continue
        sold = any(str(ws.cell(r, c).value or '').strip().lower() == 'vendido' for c in range(6, 13))
        lots.append({'mz': str(mz), 'n': str(lote), 'm2': None if sold else num(ws.cell(r, 8).value),
                     'status': 'vendido' if sold else 'disponible'})
    return lots


def paraqvaria(ws):
    lots = []
    for r in range(3, ws.max_row + 1):
        mz, tipo, lote = ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value
        if mz is None or lote is None:
            continue
        estado = str(ws.cell(r, 1).value or '').strip().lower()
        status = {'vendido': 'vendido', '': 'disponible'}.get(estado, 'nodisponible')
        lots.append({'mz': str(mz), 'n': str(lote), 'm2': num(ws.cell(r, 6).value),
                     'type': str(tipo).strip().capitalize(), 'status': status})
    return lots


def ribera(ws):
    # La hoja lista solo lotes disponibles. Los grupos de filas se separan cuando la
    # numeración de lote reinicia; la letra de manzana está en la columna A del grupo.
    groups, cur, prev = [], [], None
    for r in range(3, ws.max_row + 1):
        lote = ws.cell(r, 2).value
        if isinstance(lote, (int, float)) and ws.cell(r, 3).value:
            if prev is not None and lote < prev and cur:
                groups.append(cur)
                cur = []
            cur.append(r)
            prev = lote
    groups.append(cur)
    labels = {r: str(ws.cell(r, 1).value) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value}
    lots = []
    for rows in groups:
        mz = next(labels[r] for r in rows if r in labels)
        for r in rows:
            lots.append({'mz': mz, 'n': str(int(ws.cell(r, 2).value)), 'm2': num(ws.cell(r, 3).value),
                         'type': 'Residencial', 'status': 'disponible'})
    # Lotes de alta densidad: manzanas M1–M3, con código de lote (M1, MT1…) en la col. B.
    mz = None
    for r in range(3, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str) and a[:1] == 'M' and a[1:].isdigit():
            mz = a
        if mz and isinstance(b, str) and b[:1] == 'M':
            lots.append({'mz': mz, 'n': b, 'm2': num(ws.cell(r, 3).value), 'type': 'Alta densidad', 'status': 'disponible'})
    return lots


def sol(ws):
    units = []
    for r in range(6, ws.max_row + 1):
        unit, piso, tipo = ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value
        if not isinstance(unit, (int, float)) or not tipo:
            continue
        estado = str(ws.cell(r, 7).value or '').strip().lower()
        ubic = str(ws.cell(r, 6).value or '').strip().lower()
        units.append({'n': str(int(unit)), 'piso': int(piso), 'type': ' '.join(str(tipo).split()).replace('Dormitorios', 'dormitorios'),
                      'm2': num(ws.cell(r, 5).value), 'ubic': {'adelante': 'Frente', 'atrás': 'Contrafrente'}.get(ubic, ubic.capitalize()),
                      'status': 'vendido' if estado == 'vendido' else 'disponible'})
    return units


def php(v, ind=0):
    pad = '    ' * ind
    if isinstance(v, dict):
        return '[' + ', '.join(f"'{k}' => {php(x)}" for k, x in v.items()) + ']'
    if isinstance(v, list):
        return '[\n' + ''.join(f'{pad}    {php(x, ind + 1)},\n' for x in v) + pad + ']'
    if v is None:
        return 'null'
    if isinstance(v, (int, float)):
        return repr(v)
    return "'" + str(v).replace('\\', '\\\\').replace("'", "\\'") + "'"


wb = openpyxl.load_workbook(SRC, data_only=True)
data = {
    'cumbres': cumbres(wb['cumbres']),
    'paraqvaria': paraqvaria(wb['paraqvaria']),
    'la-ribera': ribera(wb['ribera']),
    'sol-city': sol(wb['sol city']),
}
OUT.write_text('<?php\n// Generado por tools/importar_lotes.py desde lotes.xlsx. No editar a mano.\n\nconst LOTS = '
               + php(data) + ';\n', encoding='utf-8')
for k, v in data.items():
    by = {}
    for x in v:
        by[x['status']] = by.get(x['status'], 0) + 1
    print(f'{k:11} {len(v):4} {by}')
